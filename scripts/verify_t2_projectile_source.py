"""Inventory source-tracked projectile measurements without fitting a model."""

from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime, timezone
from hashlib import sha256
from importlib.metadata import version
import json
from math import isfinite
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
CONTRACT = ROOT / "docs/T2_PROJECTILE_SOURCE_CONTRACT.json"
LOCAL = ROOT / "data/references/t2_projectile_wadsworth_2025"
AUDIT = ROOT / "artifacts/t2-projectile-source-audit.json"
SOURCE_FILES = (Path(__file__).resolve(), CONTRACT)


def _hash_bytes(data: bytes) -> str:
    return sha256(data).hexdigest()


def _hash(path: Path) -> str:
    return _hash_bytes(path.read_bytes())


def _trial_hash(rows: list[tuple[Any, ...]]) -> str:
    return _hash_bytes(json.dumps(rows, separators=(",", ":"),
                                  allow_nan=False).encode("utf-8"))


def _check_observation_rows(rows: list[tuple[Any, ...]], contract: dict) -> dict:
    expected = contract["files"]["measured_trajectories"]
    if len(rows) != expected["rows"]:
        raise ValueError("T2 measured trajectory row count differs")
    groups: dict[int, list[tuple[Any, ...]]] = defaultdict(list)
    for row in rows:
        if (len(row) != 6 or isinstance(row[0], bool) or not isinstance(row[0], int)
                or any(isinstance(value, bool) or not isinstance(value, (int, float))
                       or not isfinite(float(value)) for value in row[1:])):
            raise ValueError("T2 measured trajectory contains missing or nonfinite values")
        groups[row[0]].append(row)
    ids = sorted(groups)
    if (ids != list(range(expected["trial_ids_first"], expected["trial_ids_last"] + 1))
            or len(ids) != expected["distinct_trials"]):
        raise ValueError("T2 measured trajectory trial identity differs")
    trials = []
    for trial_id in ids:
        group = groups[trial_id]
        times = [float(row[1]) for row in group]
        if (len(group) < 3 or any(right <= left for left, right in zip(times, times[1:]))
                or len({float(row[4]) for row in group}) != 1
                or len({float(row[5]) for row in group}) != 1):
            raise ValueError(f"T2 trial time order or declared inputs differ: {trial_id}")
        trials.append({
            "trial_id": trial_id, "samples": len(group),
            "angle_degrees": float(group[0][4]),
            "declared_v0_m_s": float(group[0][5]),
            "time_min_s": min(times), "time_max_s": max(times),
            "trajectory_sha256": _trial_hash(group),
        })
    if (sum(row["samples"] for row in trials) != expected["rows"]
            or len({row["angle_degrees"] for row in trials}) != 7
            or len({row["declared_v0_m_s"] for row in trials}) != 12):
        raise ValueError("T2 measured trajectory grid differs")
    return {
        "sample_count": len(rows), "trial_count": len(trials),
        "samples_per_trial_min": min(row["samples"] for row in trials),
        "samples_per_trial_max": max(row["samples"] for row in trials),
        "distinct_angles": 7, "distinct_declared_speeds": 12,
        "observed_declared_speed_above_paper_bound_trials": [
            row["trial_id"] for row in trials
            if row["declared_v0_m_s"] > contract["article"]["paper_stated_launch_speed_upper_bound_m_s"]
        ],
        "trials": trials,
    }


def _omission_controls(rows: list[tuple[Any, ...]], contract: dict) -> dict[str, bool]:
    copies = {
        "missing_sample_rejected": rows[:-1],
        "nonfinite_measurement_rejected": [(*rows[0][:2], float("nan"), *rows[0][3:]), *rows[1:]],
        "duplicate_time_rejected": [rows[0], (rows[1][0], rows[0][1], *rows[1][2:]), *rows[2:]],
        "changed_within_trial_input_rejected": [rows[0], (*rows[1][:4], rows[1][4] + 1, rows[1][5]), *rows[2:]],
    }
    result = {}
    for name, candidate in copies.items():
        try:
            _check_observation_rows(candidate, contract)
        except ValueError:
            result[name] = True
        else:
            raise ValueError(f"T2 measured trajectory omission control was admitted: {name}")
    return result


def _source_file(path: Path) -> dict:
    return {"path": path.relative_to(ROOT).as_posix(),
            "sha256": _hash(path), "bytes": path.stat().st_size}


def evaluate() -> dict:
    contract = json.loads(CONTRACT.read_text(encoding="utf-8"))
    article = contract["article"]
    files = contract["files"]
    boundaries = contract["boundaries"]
    if (contract.get("schema_version") != "t2-projectile-source-contract-v1"
            or article.get("doi") != "10.1088/1361-6552/add2c5"
            or article.get("paper_reported_experiments") != 82
            or article.get("paper_stated_launch_speed_upper_bound_m_s") != 6.0
            or article.get("supplement_redistribution_rights_confirmed") is not False
            or boundaries != {
                "source_tracked_real_measurements": True,
                "complete_reported_experiment_set": False,
                "supplement_rights_reviewed": False,
                "velocity_column_semantics_reviewed": False,
                "physical_model_validated": False,
                "causal_effect_identified": False,
                "scientific_holdout": False,
                "research_candidate": False,
                "publication_ready": False,
                "claim_status": "unverified",
            }
            or version("openpyxl") != contract["reader"]["xlsx_version"]
            or version("pypdf") != contract["reader"]["pdf_version"]):
        raise ValueError("T2 projectile source identity, reader, or boundary differs")
    local_paths = {key: LOCAL / item["local_name"] for key, item in files.items()}
    for key, path in local_paths.items():
        if (not path.is_file() or path.stat().st_size != files[key]["bytes"]
                or _hash(path) != files[key]["sha256"]):
            raise ValueError(f"T2 projectile source file differs: {key}")
    pdf_pages = [page.extract_text() for page in PdfReader(str(local_paths["article_pdf"])).pages]
    combined = "".join(pdf_pages)
    compact = re.sub(r"\s+", "", combined).lower()
    if (len(pdf_pages) != files["article_pdf"]["pages"]
            or "creativecommonsattribution4.0licence" not in compact
            or "weperformed82experiments" not in compact
            or "lightgateaccuratetowithin0.1" not in compact):
        raise ValueError("T2 article license, experiment count, or measurement method differs")
    measurements = local_paths["measured_trajectories"]
    with ZipFile(measurements) as archive:
        if any("externallink" in name.lower() or "vbaproject" in name.lower()
               for name in archive.namelist()):
            raise ValueError("T2 measured workbook has external links or macros")
    book = load_workbook(measurements, read_only=True, data_only=False, keep_links=False)
    try:
        if book.sheetnames != [files["measured_trajectories"]["sheet"]]:
            raise ValueError("T2 measured workbook sheet differs")
        sheet = book.active
        header, *rows = list(sheet.values)
        if (list(header) != files["measured_trajectories"]["columns"]
                or sheet.max_column != 6
                or any(isinstance(value, str) and value.startswith("=")
                       for row in rows for value in row)):
            raise ValueError("T2 measured workbook columns or formula boundary differs")
        inventory = _check_observation_rows(rows, contract)
        controls = _omission_controls(rows, contract)
    finally:
        book.close()
    numerical = local_paths["numerical_spreadsheet"]
    with ZipFile(numerical) as archive:
        if any("externallink" in name.lower() or "vbaproject" in name.lower()
               for name in archive.namelist()):
            raise ValueError("T2 numerical workbook has external links or macros")
    book = load_workbook(numerical, read_only=True, data_only=False, keep_links=False)
    try:
        if book.sheetnames != [files["numerical_spreadsheet"]["sheet"]]:
            raise ValueError("T2 numerical workbook sheet differs")
        sheet = book.active
        formulas = sum(isinstance(value, str) and value.startswith("=")
                       for row in sheet.values for value in row)
        if (sheet.max_row != files["numerical_spreadsheet"]["rows"]
                or sheet.max_column != files["numerical_spreadsheet"]["columns"]
                or formulas != files["numerical_spreadsheet"]["formula_cells"]
                or files["numerical_spreadsheet"]["admitted_as_observations"] is not False):
            raise ValueError("T2 numerical workbook formula or observation boundary differs")
    finally:
        book.close()
    if (inventory["trial_count"] != 30
            or inventory["observed_declared_speed_above_paper_bound_trials"] !=
            [4, 5, 6, 10, 11, 12, 13, 14, 15, 16, 17, 18, 22, 23, 24]):
        raise ValueError("T2 trial coverage or declared speed discrepancy differs")
    return {
        "schema_version": "t2-projectile-source-audit-v1",
        "status": "verified-30-trial-real-source-inventory-only",
        "contract_sha256": _hash(CONTRACT),
        "article": article,
        "local_file_fingerprints": [
            {"role": key, "filename": path.name, "sha256": _hash(path),
             "bytes": path.stat().st_size, "redistributed": False}
            for key, path in local_paths.items()
        ],
        "article_page_text_sha256": [_hash_bytes(page.encode("utf-8")) for page in pdf_pages],
        "reader": contract["reader"],
        "measured_inventory": inventory,
        "numerical_spreadsheet": {
            "formula_cells": formulas, "admitted_as_observations": False,
        },
        "negative_controls": controls,
        "boundaries": boundaries,
        "source_files": [_source_file(path) for path in SOURCE_FILES],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    action = parser.add_mutually_exclusive_group(required=True)
    action.add_argument("--write", action="store_true")
    action.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    current = evaluate()
    if args.write:
        if AUDIT.exists():
            raise FileExistsError(f"refusing to overwrite T2 source audit: {AUDIT}")
        current["recorded_at"] = datetime.now(timezone.utc).isoformat()
        AUDIT.write_text(json.dumps(current, indent=2, sort_keys=True) + "\n",
                         encoding="utf-8", newline="\n")
    else:
        saved = json.loads(AUDIT.read_text(encoding="utf-8"))
        stamp = saved.pop("recorded_at", None)
        if (not isinstance(stamp, str)
                or datetime.fromisoformat(stamp.replace("Z", "+00:00")).tzinfo is None
                or saved != current):
            raise ValueError("T2 source audit differs from pinned local measurements")
    print(json.dumps({"status": current["status"],
                      "samples": current["measured_inventory"]["sample_count"],
                      "trials": current["measured_inventory"]["trial_count"],
                      "speed_discrepancy_trials": len(current["measured_inventory"]["observed_declared_speed_above_paper_bound_trials"]),
                      "model_validated": False, "causal_effect_identified": False},
                     sort_keys=True))


if __name__ == "__main__":
    main()
